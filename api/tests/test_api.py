import asyncio
import base64
import contextlib
import hashlib
import importlib
import io
import json
import os
import tempfile
import time
import unittest
from pathlib import Path

import fitz
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation


class ApiRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix="ppt-ai-api-test-")
        os.environ["STORAGE_BACKEND"] = "local"
        os.environ["DATA_DIR"] = cls.tmp.name
        os.environ.pop("ZHIPU_API_KEY", None)
        salt = b"ppt-ai-test-salt"
        digest = hashlib.pbkdf2_hmac("sha256", b"PPT-TEST-CODE", salt, 210000)
        encode = lambda value: base64.urlsafe_b64encode(value).decode("ascii").rstrip("=")
        os.environ["INVITE_CODES_JSON"] = json.dumps({
            "invite01": f"pbkdf2_sha256$210000${encode(salt)}${encode(digest)}"
        })
        os.environ["INVITE_EXPIRES_AT"] = "2099-01-01T00:00:00Z"
        os.environ["ACCESS_SESSION_SECONDS"] = str(30 * 24 * 60 * 60)
        os.environ["ACCESS_AUTH_SECRET"] = "unit-test-auth-secret"

        import main

        cls.main = importlib.reload(main)
        # 必须作为上下文管理器进入：否则 TestClient 每个请求起一个临时事件循环，
        # 请求一返回就拆掉，analyze 起的后台任务会被连带取消，
        # 异步分析永远等不到 done。生产环境跑在常驻的 uvicorn 循环上，没这个问题。
        cls._client_ctx = TestClient(cls.main.app)
        cls.client = cls._client_ctx.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls._client_ctx.__exit__(None, None, None)
        cls.tmp.cleanup()

    def test_project_file_and_persistence_flow(self):
        health = self.client.get("/health")
        self.assertEqual(health.status_code, 200)
        self.assertTrue(health.json()["storageOk"])
        self.assertTrue(health.json()["accessControlReady"])

        denied_create = self.client.post("/projects", json={"title": "Denied"})
        self.assertEqual(denied_create.status_code, 401)

        denied_login = self.client.post(
            "/auth/login",
            json={"code": "PPT-WRONG-CODE"},
        )
        self.assertEqual(denied_login.status_code, 401)

        logged_in = self.client.post(
            "/auth/login",
            json={"code": "ppt-test-code"},
        )
        self.assertEqual(logged_in.status_code, 200)
        access_token = logged_in.json()["token"]
        access_headers = {"Authorization": f"Bearer {access_token}"}
        session = self.client.get("/auth/session", headers=access_headers)
        self.assertEqual(session.status_code, 200)
        self.assertEqual(session.json()["user"]["username"], "invite01")

        created = self.client.post(
            "/projects",
            headers=access_headers,
            json={
                "title": "PPT 测试项目",
                "purpose": "验证资料顺序锁与持久化",
                "audience": "测试人员",
                "limit": "8 页",
            },
        )
        self.assertEqual(created.status_code, 200)
        body = created.json()
        self.assertEqual(body["project"]["createdBy"], "invite01")
        project_id = body["project"]["id"]
        token = body["accessToken"]
        headers = {"X-Project-Token": token}

        unauthorized = self.client.get(f"/projects/{project_id}")
        self.assertEqual(unauthorized.status_code, 401)

        uploaded = self.client.post(
            f"/projects/{project_id}/files",
            headers=headers,
            files=[
                ("files", ("第一份.txt", "第一份资料：市场规模 100。".encode("utf-8"), "text/plain")),
                ("files", ("第二份.txt", "第二份资料：客户满意度 95%。".encode("utf-8"), "text/plain")),
            ],
        )
        self.assertEqual(uploaded.status_code, 200)
        materials = uploaded.json()["project"]["materials"]
        self.assertEqual(len(materials), 2)

        first_id = materials[0]["id"]
        second_id = materials[1]["id"]
        patched = self.client.patch(
            f"/projects/{project_id}",
            headers=headers,
            json={
                "locked": True,
                "materials": [
                    {"id": first_id, "customerOrder": 2},
                    {"id": second_id, "customerOrder": 1},
                ],
            },
        )
        self.assertEqual(patched.status_code, 200)
        patched_project = patched.json()["project"]
        self.assertTrue(patched_project["locked"])
        self.assertEqual([m["id"] for m in patched_project["materials"]], [second_id, first_id])

        downloaded = self.client.get(
            f"/projects/{project_id}/files/{second_id}", headers=headers
        )
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn("客户满意度", downloaded.content.decode("utf-8"))

        self.main.storage = self.main.LocalStorage(Path(self.tmp.name))
        restored = self.client.get(f"/projects/{project_id}", headers=headers)
        self.assertEqual(restored.status_code, 200)
        self.assertTrue(restored.json()["project"]["locked"])
        self.assertEqual(len(restored.json()["project"]["materials"]), 2)

        no_key = self.client.post(f"/projects/{project_id}/analyze", headers=headers)
        self.assertEqual(no_key.status_code, 503)

    def test_supported_document_parsers(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "PDF parser test")
        pdf_text, previews = self.main.parse_pdf(pdf.tobytes())
        self.assertIn("PDF parser test", pdf_text)
        self.assertEqual(len(previews), 1)

        docx_buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("DOCX parser test")
        document.save(docx_buffer)
        self.assertIn("DOCX parser test", self.main.parse_docx(docx_buffer.getvalue()))

        xlsx_buffer = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "XLSX parser test"
        workbook.save(xlsx_buffer)
        self.assertIn("XLSX parser test", self.main.parse_xlsx(xlsx_buffer.getvalue()))

        pptx_buffer = io.BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "PPTX parser test"
        presentation.save(pptx_buffer)
        self.assertIn("PPTX parser test", self.main.parse_pptx(pptx_buffer.getvalue()))

    def test_access_token_never_outlives_invite(self):
        original_expiry = self.main.INVITE_EXPIRES_AT_TIMESTAMP
        try:
            self.main.INVITE_EXPIRES_AT_TIMESTAMP = int(time.time()) + 60
            token, expires_at = self.main.issue_access_token("invite01")
            encoded = token.split(".", 1)[0]
            payload = json.loads(self.main.b64url_decode(encoded).decode("utf-8"))
            self.assertEqual(payload["exp"], self.main.INVITE_EXPIRES_AT_TIMESTAMP)
            self.assertEqual(
                expires_at,
                self.main.datetime.fromtimestamp(
                    self.main.INVITE_EXPIRES_AT_TIMESTAMP,
                    tz=self.main.timezone.utc,
                ).isoformat(),
            )
        finally:
            self.main.INVITE_EXPIRES_AT_TIMESTAMP = original_expiry

    def test_ai_retries_transient_provider_overload(self):
        class FakeResponse:
            def __init__(self, status_code, payload, headers=None):
                self.status_code = status_code
                self._payload = payload
                self.headers = headers or {}
                self.text = json.dumps(payload, ensure_ascii=False)

            def json(self):
                return self._payload

        class FakeClient:
            calls = 0

            def __init__(self, timeout):
                self.timeout = timeout

            async def __aenter__(self):
                return self

            async def __aexit__(self, exc_type, exc, tb):
                return False

            async def post(self, url, headers, json):
                FakeClient.calls += 1
                if FakeClient.calls == 1:
                    return FakeResponse(429, {"error": {"code": "1305"}}, {"Retry-After": "1"})
                return FakeResponse(200, {"choices": [{"message": {"content": "ok"}}]})

        original_client = self.main.httpx.AsyncClient
        original_key = self.main.ZHIPU_API_KEY
        original_attempts = self.main.AI_MAX_ATTEMPTS
        original_sleep = self.main.asyncio.sleep
        try:
            self.main.httpx.AsyncClient = FakeClient
            self.main.ZHIPU_API_KEY = "test-key"
            self.main.AI_MAX_ATTEMPTS = 3

            async def no_sleep(_delay):
                return None

            self.main.asyncio.sleep = no_sleep
            result = asyncio.run(self.main.zhipu_chat({"model": "test"}))
            self.assertEqual(result["choices"][0]["message"]["content"], "ok")
            self.assertEqual(FakeClient.calls, 2)
        finally:
            self.main.httpx.AsyncClient = original_client
            self.main.ZHIPU_API_KEY = original_key
            self.main.AI_MAX_ATTEMPTS = original_attempts
            self.main.asyncio.sleep = original_sleep

    # ------------------------------------------------------------------
    # 分析异步化：POST 立即返回 202，结果通过 GET 轮询
    # ------------------------------------------------------------------
    def _new_project_with_material(self, title="异步分析测试"):
        logged_in = self.client.post("/auth/login", json={"code": "ppt-test-code"})
        self.assertEqual(logged_in.status_code, 200)
        access_headers = {"Authorization": f"Bearer {logged_in.json()['token']}"}
        created = self.client.post("/projects", headers=access_headers, json={"title": title})
        self.assertEqual(created.status_code, 200)
        project_id = created.json()["project"]["id"]
        headers = {"X-Project-Token": created.json()["accessToken"]}
        uploaded = self.client.post(
            f"/projects/{project_id}/files",
            headers=headers,
            files=[("files", ("素材.txt", "测试素材内容。".encode("utf-8"), "text/plain"))],
        )
        self.assertEqual(uploaded.status_code, 200)
        return project_id, headers

    @contextlib.contextmanager
    def _stubbed_ai(self, *, fail=None, delay=0.0, counter=None):
        """替换掉真实 AI 调用。run_analysis 通过模块全局引用这两个函数，所以打模块属性即可。"""
        main = self.main
        original = (main.ZHIPU_API_KEY, main.analyze_material, main.synthesize_outline)
        main.ZHIPU_API_KEY = "test-key"

        async def fake_analyze_material(item):
            if counter is not None:
                counter.append(item["id"])
            if delay:
                await asyncio.sleep(delay)
            if fail is not None:
                raise fail
            return {"id": item["id"], "name": item.get("name", ""), "summary": "测试摘要"}

        async def fake_synthesize_outline(project, analyses):
            return {
                "outline": [{"title": "测试章节", "summary": "摘要"}],
                "pageEstimate": "8页",
                "orderSuggestions": [],
            }

        main.analyze_material = fake_analyze_material
        main.synthesize_outline = fake_synthesize_outline
        try:
            yield
        finally:
            main.ZHIPU_API_KEY, main.analyze_material, main.synthesize_outline = original

    def _wait_for_status(self, project_id, headers, target, timeout=15.0):
        deadline = time.time() + timeout
        payload = None
        while time.time() < deadline:
            response = self.client.get(f"/projects/{project_id}/analyze", headers=headers)
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            if payload["analysisStatus"] == target:
                return payload
            time.sleep(0.05)
        self.fail(f"等待 analysisStatus={target} 超时，最后一次为 {payload}")

    def test_analyze_returns_immediately_and_finishes_in_background(self):
        project_id, headers = self._new_project_with_material()
        with self._stubbed_ai(delay=0.2):
            started = self.client.post(f"/projects/{project_id}/analyze", headers=headers)
            # 关键：不等 AI 跑完就返回，否则线上会撞网关超时
            self.assertEqual(started.status_code, 202)
            self.assertEqual(started.json()["analysisStatus"], "running")
            self.assertNotIn("result", started.json())

            done = self._wait_for_status(project_id, headers, "done")
            self.assertTrue(done["result"]["outline"])
            self.assertEqual(len(done["materials"]), 1)
            self.assertTrue(done["completedAt"])

        # 结果同样要落到项目上，刷新页面后才拿得到
        project = self.client.get(f"/projects/{project_id}", headers=headers).json()["project"]
        self.assertEqual(project["analysisStatus"], "done")
        self.assertTrue(project["aiResult"]["outline"])

    def test_analyze_failure_surfaces_through_status_endpoint(self):
        project_id, headers = self._new_project_with_material("分析失败")
        with self._stubbed_ai(fail=RuntimeError("AI 服务炸了")):
            started = self.client.post(f"/projects/{project_id}/analyze", headers=headers)
            # 失败发生在后台，POST 本身仍然是成功受理
            self.assertEqual(started.status_code, 202)

            failed = self._wait_for_status(project_id, headers, "failed")
            self.assertIn("AI 服务炸了", failed["error"])
            self.assertNotIn("result", failed)

    def test_duplicate_submit_does_not_start_a_second_analysis(self):
        project_id, headers = self._new_project_with_material("重复提交")
        calls: list = []
        with self._stubbed_ai(delay=0.5, counter=calls):
            first = self.client.post(f"/projects/{project_id}/analyze", headers=headers)
            second = self.client.post(f"/projects/{project_id}/analyze", headers=headers)
            self.assertEqual(first.status_code, 202)
            self.assertEqual(second.status_code, 202)
            self.assertEqual(second.json()["analysisStatus"], "running")

            self._wait_for_status(project_id, headers, "done")
            # 连点两次不该把同一批素材打给 AI 两遍
            self.assertEqual(len(calls), 1)

    def test_interrupted_run_is_reported_as_failed_and_can_restart(self):
        project_id, headers = self._new_project_with_material("中断恢复")

        # 模拟容器在分析途中重启：状态永久停在 running
        project = self.main.load_project(project_id)
        project["analysisStatus"] = "running"
        project["analysisStartedAt"] = "2020-01-01T00:00:00+00:00"
        self.main.save_project(project)

        stuck = self.client.get(f"/projects/{project_id}/analyze", headers=headers)
        self.assertEqual(stuck.json()["analysisStatus"], "failed")
        self.assertIn("中断", stuck.json()["error"])

        # 而且必须允许重新发起，不能被那个永不结束的状态锁死
        with self._stubbed_ai():
            restarted = self.client.post(f"/projects/{project_id}/analyze", headers=headers)
            self.assertEqual(restarted.status_code, 202)
            self._wait_for_status(project_id, headers, "done")


if __name__ == "__main__":
    unittest.main()
